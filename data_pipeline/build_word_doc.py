import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_word_report():
    doc = Document()

    # Page Margins (1 inch all around)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Base Colors
    COLOR_PRIMARY = RGBColor(0, 113, 227)     # Apple Blue (#0071e3)
    COLOR_DARK = RGBColor(29, 29, 31)        # Dark Charcoal (#1d1d1f)
    COLOR_MUTED = RGBColor(110, 110, 115)    # Muted Gray (#6e6e73)

    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Apple Customer Support Assistant")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY
    title_p.paragraph_format.space_after = Pt(4)

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Hiver SDE Intern Take-Home Assignment Technical Report\nAuthor: Shriya Mohanty | Repository: https://github.com/shriya-0802/Hiver-SDE-Intern-Assignment.git")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = COLOR_MUTED
    sub_p.paragraph_format.space_after = Pt(18)

    # Read README.md
    with open("README.md", "r", encoding="utf-8") as f:
        readme_text = f.read()

    sections_text = readme_text.split("\n## ")

    for i, sec in enumerate(sections_text):
        if i == 0:
            continue
        lines = sec.strip().split("\n")
        header_title = lines[0].replace("#", "").strip()

        # Section Heading 1
        h1 = doc.add_heading(level=1)
        h1_run = h1.add_run(header_title)
        h1_run.font.name = 'Arial'
        h1_run.font.size = Pt(15)
        h1_run.font.bold = True
        h1_run.font.color.rgb = COLOR_PRIMARY
        h1.paragraph_format.space_before = Pt(16)
        h1.paragraph_format.space_after = Pt(6)

        in_table = False
        table_rows = []

        for line in lines[1:]:
            line_str = line.strip()
            if not line_str or line_str == "---":
                continue

            # Table parsing
            if line_str.startswith("|") and line_str.endswith("|"):
                if "---" in line_str:
                    continue
                cols = [c.strip() for c in line_str.split("|")[1:-1]]
                table_rows.append(cols)
                in_table = True
                continue
            else:
                if in_table and table_rows:
                    # Render Word Table
                    num_rows = len(table_rows)
                    num_cols = len(table_rows[0])
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    for r_idx, row_data in enumerate(table_rows):
                        row = tbl.rows[r_idx]
                        for c_idx, cell_value in enumerate(row_data):
                            cell = row.cells[c_idx]
                            cell.text = cell_value
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(3)
                            p.paragraph_format.space_after = Pt(3)
                            
                            for run in p.runs:
                                run.font.name = 'Arial'
                                run.font.size = Pt(9)
                                if r_idx == 0:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                else:
                                    run.font.color.rgb = COLOR_DARK

                            if r_idx == 0:
                                set_cell_background(cell, "0071E3")
                            elif r_idx % 2 == 1:
                                set_cell_background(cell, "F5F5F7")
                            else:
                                set_cell_background(cell, "FFFFFF")

                    doc.add_paragraph() # Spacer after table
                    table_rows = []
                    in_table = False

            # Subheadings
            if line_str.startswith("### "):
                sub_title = line_str.replace("### ", "").strip()
                h2 = doc.add_heading(level=2)
                h2_run = h2.add_run(sub_title)
                h2_run.font.name = 'Arial'
                h2_run.font.size = Pt(12)
                h2_run.font.bold = True
                h2_run.font.color.rgb = COLOR_DARK
                h2.paragraph_format.space_before = Pt(10)
                h2.paragraph_format.space_after = Pt(4)
                continue

            # Code block lines
            if line_str.startswith("```"):
                continue

            # Bullet points
            if line_str.startswith("- ") or line_str.startswith("* ") or (len(line_str) > 2 and line_str[0].isdigit() and line_str[1] == '.'):
                bullet_p = doc.add_paragraph(style='List Bullet')
                b_text = line_str.lstrip("-*0123456789. ").strip()
                b_run = bullet_p.add_run(b_text)
                b_run.font.name = 'Arial'
                b_run.font.size = Pt(10)
                b_run.font.color.rgb = COLOR_DARK
                bullet_p.paragraph_format.space_before = Pt(1)
                bullet_p.paragraph_format.space_after = Pt(3)
                continue

            # Normal Paragraph
            p = doc.add_paragraph()
            p_run = p.add_run(line_str)
            p_run.font.name = 'Arial'
            p_run.font.size = Pt(10)
            p_run.font.color.rgb = COLOR_DARK
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)

        # Flush trailing table if any
        if in_table and table_rows:
            num_rows = len(table_rows)
            num_cols = len(table_rows[0])
            tbl = doc.add_table(rows=num_rows, cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row_data in enumerate(table_rows):
                row = tbl.rows[r_idx]
                for c_idx, cell_value in enumerate(row_data):
                    cell = row.cells[c_idx]
                    cell.text = cell_value
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(9)
                        if r_idx == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            run.font.color.rgb = COLOR_DARK

                    if r_idx == 0:
                        set_cell_background(cell, "0071E3")
                    elif r_idx % 2 == 1:
                        set_cell_background(cell, "F5F5F7")
                    else:
                        set_cell_background(cell, "FFFFFF")

            doc.add_paragraph()
            table_rows = []
            in_table = False

    docx_path = "Apple_Customer_Support_Assistant_Report.docx"
    doc.save(docx_path)
    print(f"Successfully generated Word Document: {docx_path}")
    return docx_path

if __name__ == "__main__":
    create_word_report()
