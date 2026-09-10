import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def build_all_reports():
    docx_path = "Apple_Customer_Support_Assistant_Report.docx"
    pdf_path = "Apple_Customer_Support_Assistant_Report.pdf"

    # Read README.md
    with open("README.md", "r", encoding="utf-8") as f:
        readme_text = f.read()

    # 1. BUILD WORD DOCUMENT (.docx)
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 113, 227)
    COLOR_DARK = RGBColor(29, 29, 31)
    COLOR_MUTED = RGBColor(110, 110, 115)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Apple Customer Support Assistant")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    p_title.paragraph_format.space_after = Pt(4)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Hiver SDE Intern Take-Home Assignment Technical Report\nAuthor: Shriya Mohanty | Repository: https://github.com/shriya-0802/Hiver-SDE-Intern-Assignment.git")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_MUTED
    p_sub.paragraph_format.space_after = Pt(16)

    sections_text = readme_text.split("\n## ")

    for i, sec in enumerate(sections_text):
        if i == 0:
            continue
        lines = sec.strip().split("\n")
        header_title = lines[0].replace("#", "").strip()

        h1 = doc.add_heading(level=1)
        h1_run = h1.add_run(header_title)
        h1_run.font.name = 'Arial'
        h1_run.font.size = Pt(14)
        h1_run.font.bold = True
        h1_run.font.color.rgb = COLOR_PRIMARY
        h1.paragraph_format.space_before = Pt(16)
        h1.paragraph_format.space_after = Pt(6)

        in_table = False
        table_rows = []

        for line in lines[1:]:
            line_str = line.strip()
            if not line_str or line_str == "---":
                continue

            if line_str.startswith("|") and line_str.endswith("|"):
                if "---" in line_str:
                    continue
                cols = [c.strip() for c in line_str.split("|")[1:-1]]
                table_rows.append(cols)
                in_table = True
                continue
            else:
                if in_table and table_rows:
                    num_rows = len(table_rows)
                    num_cols = len(table_rows[0])
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    for r_idx, row_data in enumerate(table_rows):
                        row = tbl.rows[r_idx]
                        for c_idx, cell_value in enumerate(row_data):
                            cell = row.cells[c_idx]
                            cell.text = cell_value
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(4)
                            p.paragraph_format.space_after = Pt(4)
                            
                            for run in p.runs:
                                run.font.name = 'Arial'
                                run.font.size = Pt(9.5)
                                if r_idx == 0:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                else:
                                    run.font.color.rgb = COLOR_DARK

                            if r_idx == 0:
                                set_cell_background(cell, "0071E3")
                            elif r_idx % 2 == 1:
                                set_cell_background(cell, "F5F5F7")
                            else:
                                set_cell_background(cell, "FFFFFF")

                    doc.add_paragraph()
                    table_rows = []
                    in_table = False

            if line_str.startswith("### "):
                sub_title = line_str.replace("### ", "").strip()
                h2 = doc.add_heading(level=2)
                h2_run = h2.add_run(sub_title)
                h2_run.font.name = 'Arial'
                h2_run.font.size = Pt(12)
                h2_run.font.bold = True
                h2_run.font.color.rgb = COLOR_DARK
                h2.paragraph_format.space_before = Pt(12)
                h2.paragraph_format.space_after = Pt(4)
                continue

            if line_str.startswith("```"):
                continue

            if line_str.startswith("- ") or line_str.startswith("* ") or (len(line_str) > 2 and line_str[0].isdigit() and line_str[1] == '.'):
                bullet_p = doc.add_paragraph(style='List Bullet')
                b_text = line_str.lstrip("-*0123456789. ").strip()
                b_run = bullet_p.add_run(b_text)
                b_run.font.name = 'Arial'
                b_run.font.size = Pt(10)
                b_run.font.color.rgb = COLOR_DARK
                bullet_p.paragraph_format.space_before = Pt(2)
                bullet_p.paragraph_format.space_after = Pt(4)
                continue

            p = doc.add_paragraph()
            p_run = p.add_run(line_str)
            p_run.font.name = 'Arial'
            p_run.font.size = Pt(10)
            p_run.font.color.rgb = COLOR_DARK
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)

        if in_table and table_rows:
            num_rows = len(table_rows)
            num_cols = len(table_rows[0])
            tbl = doc.add_table(rows=num_rows, cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row_data in enumerate(table_rows):
                row = tbl.rows[r_idx]
                for c_idx, cell_value in enumerate(row_data):
                    cell = row.cells[c_idx]
                    cell.text = cell_value
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    for run in p.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(9.5)
                        if r_idx == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            run.font.color.rgb = COLOR_DARK

                    if r_idx == 0:
                        set_cell_background(cell, "0071E3")
                    elif r_idx % 2 == 1:
                        set_cell_background(cell, "F5F5F7")
                    else:
                        set_cell_background(cell, "FFFFFF")

            doc.add_paragraph()
            table_rows = []
            in_table = False

    doc.save(docx_path)
    print(f"Successfully generated Word Document: {docx_path}")

    # 2. BUILD HIGH QUALITY 7-8 PAGE PDF USING REPORTLAB
    pdf_doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    PRIMARY_PDF = colors.HexColor("#0071e3")
    DARK_PDF = colors.HexColor("#1d1d1f")
    MUTED_PDF = colors.HexColor("#6e6e73")
    BG_LIGHT_PDF = colors.HexColor("#f5f5f7")
    BORDER_PDF = colors.HexColor("#d2d2d7")

    title_style = ParagraphStyle('PdfTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=22, leading=26, textColor=PRIMARY_PDF, spaceAfter=4)
    sub_style = ParagraphStyle('PdfSub', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=10.5, leading=14, textColor=MUTED_PDF, spaceAfter=14)
    h1_style = ParagraphStyle('PdfH1', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=14, leading=18, textColor=PRIMARY_PDF, spaceBefore=14, spaceAfter=8, keepWithNext=True)
    h2_style = ParagraphStyle('PdfH2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11.5, leading=15, textColor=DARK_PDF, spaceBefore=10, spaceAfter=6, keepWithNext=True)
    body_style = ParagraphStyle('PdfBody', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14.5, textColor=DARK_PDF, spaceAfter=6)
    bullet_style = ParagraphStyle('PdfBullet', parent=body_style, leftIndent=14, firstLineIndent=-10, spaceAfter=4)
    th_style = ParagraphStyle('PdfTh', fontName='Helvetica-Bold', fontSize=9, leading=11, textColor=colors.white)
    tc_style = ParagraphStyle('PdfTc', fontName='Helvetica', fontSize=8.5, leading=11, textColor=DARK_PDF)

    pdf_story = []
    pdf_story.append(Paragraph("Apple Customer Support Assistant", title_style))
    pdf_story.append(Paragraph("<b>Hiver SDE Intern Take-Home Assignment Technical Report</b><br/>Author: Shriya Mohanty | Repository: https://github.com/shriya-0802/Hiver-SDE-Intern-Assignment.git", sub_style))
    pdf_story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY_PDF, spaceAfter=12))

    for i, sec in enumerate(sections_text):
        if i == 0:
            continue
        lines = sec.strip().split("\n")
        header_title = lines[0].replace("#", "").strip()

        pdf_story.append(Paragraph(f"## {header_title}", h1_style))
        pdf_story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_PDF, spaceAfter=8))

        in_table = False
        table_rows = []

        for line in lines[1:]:
            line_str = line.strip()
            if not line_str or line_str == "---":
                continue

            if line_str.startswith("|") and line_str.endswith("|"):
                if "---" in line_str:
                    continue
                cols = [c.strip() for c in line_str.split("|")[1:-1]]
                table_rows.append(cols)
                in_table = True
                continue
            else:
                if in_table and table_rows and len(table_rows[0]) > 0:
                    col_count = len(table_rows[0])
                    col_width = 520 / max(1, col_count)
                    table_data = []
                    for r_idx, row in enumerate(table_rows):
                        formatted_row = []
                        for cell in row:
                            cell_p = Paragraph(cell, th_style if r_idx == 0 else tc_style)
                            formatted_row.append(cell_p)
                        table_data.append(formatted_row)
                    
                    t = Table(table_data, colWidths=[col_width]*col_count)
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0,0), (-1,0), PRIMARY_PDF),
                        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                        ('GRID', (0,0), (-1,-1), 0.5, BORDER_PDF),
                        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, BG_LIGHT_PDF]),
                        ('TOPPADDING', (0,0), (-1,-1), 6),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                    ]))
                    pdf_story.append(t)
                    pdf_story.append(Spacer(1, 10))
                    table_rows = []
                    in_table = False

            if line_str.startswith("### "):
                sub_title = line_str.replace("### ", "").strip()
                pdf_story.append(Paragraph(sub_title, h2_style))
                continue

            if line_str.startswith("```"):
                continue

            if line_str.startswith("- ") or line_str.startswith("* ") or (len(line_str) > 2 and line_str[0].isdigit() and line_str[1] == '.'):
                item_text = line_str.lstrip("-*0123456789. ").strip()
                pdf_story.append(Paragraph(f"• {item_text}", bullet_style))
                continue

            pdf_story.append(Paragraph(line_str, body_style))

        if in_table and table_rows and len(table_rows[0]) > 0:
            col_count = len(table_rows[0])
            col_width = 520 / max(1, col_count)
            table_data = []
            for r_idx, row in enumerate(table_rows):
                formatted_row = []
                for cell in row:
                    cell_p = Paragraph(cell, th_style if r_idx == 0 else tc_style)
                    formatted_row.append(cell_p)
                table_data.append(formatted_row)
            
            t = Table(table_data, colWidths=[col_width]*col_count)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), PRIMARY_PDF),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('GRID', (0,0), (-1,-1), 0.5, BORDER_PDF),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, BG_LIGHT_PDF]),
                ('TOPPADDING', (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ]))
            pdf_story.append(t)
            pdf_story.append(Spacer(1, 10))

        pdf_story.append(Spacer(1, 12))

    def add_header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(MUTED_PDF)
        canvas.drawString(45, 20, "Apple Customer Support Assistant — Hiver SDE Intern Technical Report")
        canvas.drawRightString(567, 20, f"Page {doc.page}")
        canvas.restoreState()

    pdf_doc.build(pdf_story, onFirstPage=add_header_footer, onLaterPages=add_header_footer)
    print(f"Successfully generated PDF Document: {pdf_path}")

if __name__ == "__main__":
    build_all_reports()
